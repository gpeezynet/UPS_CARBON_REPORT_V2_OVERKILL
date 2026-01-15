"""Tests for the Miles fill helper script."""

import openpyxl

from tools.fill_miles_excel import fill_miles_in_workbook


def _write_sample_workbook(path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Outbound"
    sheet.append(["Reference No.1", "Receiver Zip Code", "Miles"])
    sheet.append(["SO-1", "90210", None])
    sheet.append(["SO-2", "10001", ""])
    sheet.append(["SO-3", "", None])
    workbook.save(path)


def test_fill_miles_from_zip(tmp_path):
    input_path = tmp_path / "input.xlsx"
    output_path = tmp_path / "output.xlsx"
    _write_sample_workbook(input_path)

    summaries = fill_miles_in_workbook(input_path, output_path, origin_zip="35763")

    assert output_path.exists()
    assert summaries[0].total_rows == 3

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook["Outbound"]

    miles_row_1 = sheet.cell(row=2, column=3).value
    miles_row_2 = sheet.cell(row=3, column=3).value
    miles_row_3 = sheet.cell(row=4, column=3).value

    assert isinstance(miles_row_1, (int, float)) and miles_row_1 > 0
    assert isinstance(miles_row_2, (int, float)) and miles_row_2 > 0
    assert miles_row_3 in (None, "")
